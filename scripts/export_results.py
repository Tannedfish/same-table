"""Write artifacts/Experimental_results.xlsx from the two notebooks' result files.

    python scripts/export_results.py

Reads artifacts/p1_results.json (Model1 notebook) and artifacts/experiments.json (Model2 notebook).
Accuracy cells are Excel formulas over the per-sample sheets, so they can be audited.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "artifacts"

FONT = "Arial"
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=12)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4D3A")
PICK_FILL = PatternFill("solid", fgColor="E3F2D0")
INPUT_FILL = PatternFill("solid", fgColor="FFF4D6")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def put(ws, ref, value, font=None, fill=None, fmt=None, box=False, wrap=False):
    cell = ws[ref]
    cell.value = value
    cell.font = font or Font(name=FONT)
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if box:
        cell.border = BOX
    if wrap:
        cell.alignment = WRAP
    return cell


def header(ws, row, col_values):
    for col, text in col_values.items():
        put(ws, f"{col}{row}", text, font=HEAD, fill=HEAD_FILL, box=True, wrap=True)


def main() -> None:
    p1 = json.loads((OUT / "p1_results.json").read_text(encoding="utf-8"))
    ex = json.loads((OUT / "experiments.json").read_text(encoding="utf-8"))
    wb = Workbook()

    # ------------------------------------------------------------ P1 test sheet
    t1 = wb.active
    t1.title = "P1 test"
    header(t1, 1, {"A": "Hand-written test sentence", "B": "Expected tags", "C": f"Predicted tags ({p1['chosen']}, fine-tuned)", "D": "Correct or not"})
    for i, row in enumerate(p1["test_detail"], start=2):
        put(t1, f"A{i}", row["sentence"], wrap=True)
        put(t1, f"B{i}", row["expected"])
        put(t1, f"C{i}", row["predicted"])
        put(t1, f"D{i}", f'=IF(B{i}=C{i},"Correct","Incorrect")')
    n1 = len(p1["test_detail"]) + 1
    for col, width in zip("ABCD", (58, 30, 34, 14)):
        t1.column_dimensions[col].width = width
    t1.freeze_panes = "A2"

    # ---------------------------------------------------------- P2 queries sheet
    t2 = wb.create_sheet("P2 queries")
    systems = list(ex["p2_ranks"])  # TF-IDF first, then the three encoders
    header(t2, 1, {"A": "Query", "B": "Target dish", **{get_column_letter(3 + k): f"Rank of target: {name}" for k, name in enumerate(systems)}})
    for i, q in enumerate(ex["p2_queries"], start=2):
        put(t2, f"A{i}", q["query"])
        put(t2, f"B{i}", q["target"])
        for k, name in enumerate(systems):
            put(t2, f"{get_column_letter(3 + k)}{i}", ex["p2_ranks"][name][i - 2])
    n2 = len(ex["p2_queries"]) + 1
    t2.column_dimensions["A"].width = 44
    t2.column_dimensions["B"].width = 40
    for k in range(len(systems)):
        t2.column_dimensions[get_column_letter(3 + k)].width = 24
    t2.row_dimensions[1].height = 48
    t2.freeze_panes = "A2"
    rank_col = {name: get_column_letter(3 + k) for k, name in enumerate(systems)}

    # ---------------------------------------------------------- Performance sheet
    ws = wb.create_sheet("Performance", 0)
    put(ws, "C1", "Project:", font=BOLD)
    put(ws, "D1", "Same Table / 糖友同枱 - shared home dinner planning for diabetes patients (Hospital Authority, classroom prototype)", font=TITLE)
    put(ws, "D2", "There are two pipelines in this app.")
    put(ws, "D3", "Pipeline 1 (understand the cook's request) is highly relevant to the project objective and therefore, it requires finetuning.")
    put(ws, "D4", "Pipeline 2 (retrieve dishes from the labeled library) does not require finetuning.")
    put(ws, "C6", "Experiment: Selecting the optimal model(s)", font=TITLE)
    put(ws, "J6", "Experiment: Assessing the overall performance", font=TITLE)

    # Pipeline 1
    put(ws, "C8", "Step 1: Select THREE relevant and comparable pre-trained models", font=BOLD)
    put(ws, "E8", "Step 2a: Finetune the selected models with own dataset", font=BOLD)
    put(ws, "C9", "Pipeline 1 - Multi-label text classification (to be finetuned)", font=BOLD)
    put(ws, "E9", "Fine-tuning")
    put(ws, "E10", "Dataset:")
    put(ws, "F10", "data/train/p1_train.jsonl, p1_eval.jsonl, p1_test.jsonl")
    put(ws, "E11", "training samples:")
    put(ws, "F11", p1["dataset"]["training"])
    put(ws, "E12", "evaluation samples:")
    put(ws, "F12", p1["dataset"]["evaluation"])
    put(ws, "E13", "testing samples (hand-written):")
    put(ws, "F13", "=COUNTA('P1 test'!A2:A%d)" % n1)
    header(ws, 15, {"D": "Model name", "E": "Accuracy (after finetuning, using trainer.evaluate())",
                    "F": "Run time (in sec. per 100 requests, CPU)", "G": "Fine-tuning time (in sec.)",
                    "H": "Accuracy on hand-written test"})
    row = 16
    for label, model in p1["candidates"].items():
        r = p1["results"][model]
        fill = PICK_FILL if model == p1["chosen"] else None
        put(ws, f"C{row}", label, box=True, fill=fill)
        put(ws, f"D{row}", model, box=True, fill=fill)
        put(ws, f"E{row}", round(r["eval_accuracy"], 4), fmt="0.000", box=True, fill=fill)
        put(ws, f"F{row}", round(r["cpu_seconds_per_100_requests"], 2), fmt="0.00", box=True, fill=fill)
        put(ws, f"G{row}", round(r["train_seconds"], 1), fmt="0", box=True, fill=fill)
        if model == p1["chosen"]:
            put(ws, f"H{row}", "=COUNTIF('P1 test'!D2:D%d,\"Correct\")/COUNTA('P1 test'!D2:D%d)" % (n1, n1), fmt="0.000", box=True, fill=fill)
        else:
            put(ws, f"H{row}", round(r["test_accuracy"], 4), fmt="0.000", box=True)
        row += 1
    kw = p1["baselines"]["keyword lexicon (no model)"]
    put(ws, f"C{row}", "Baseline", box=True)
    put(ws, f"D{row}", "keyword lexicon (no model, no training)", box=True)
    put(ws, f"E{row}", round(kw["eval_accuracy"], 4), fmt="0.000", box=True)
    put(ws, f"F{row}", "< 0.01", box=True)
    put(ws, f"G{row}", 0, box=True)
    put(ws, f"H{row}", round(kw["test_accuracy"], 4), fmt="0.000", box=True)
    row += 1
    zs = p1["baselines"].get("zero-shot mDeBERTa (no training)")
    if zs:
        put(ws, f"C{row}", "Baseline", box=True)
        put(ws, f"D{row}", "MoritzLaurer/mDeBERTa-v3-base-mnli-xnli (zero-shot, no training)", box=True)
        put(ws, f"E{row}", "n/a", box=True)
        put(ws, f"F{row}", round(100 * zs["test_seconds"] / p1["dataset"]["testing"], 2), fmt="0.00", box=True)
        put(ws, f"G{row}", 0, box=True)
        put(ws, f"H{row}", round(zs["test_accuracy"], 4), fmt="0.000", box=True)
        row += 1
    s = p1["settings"]
    put(ws, f"D{row}", f"Accuracy = all 8 tags of a sentence correct (exact match). Settings: {s['epochs']} epochs, learning rate {s['learning_rate']}, "
                       f"batch {s['batch_size']}, fine-tuned on {s['device']}. Source: notebooks/Model1_finetune_tagger.ipynb. "
                       "The picked model's hand-written accuracy is a formula over sheet 'P1 test'.", wrap=True)
    ws.merge_cells(f"D{row}:H{row}")
    ws.row_dimensions[row].height = 44
    row += 2
    put(ws, f"E{row}", "Step 2b: Pick the BEST finetuned model for implementation", font=BOLD)
    put(ws, f"E{row + 1}", f"{p1['chosen']} is picked: highest evaluation accuracy; models within one standard error of the best are tied and the shortest run time wins.")
    put(ws, f"E{row + 2}", f"The finetuned model is uploaded to Hugging Face as {ex['p1_model']} and is used for implementing the App.")
    row += 5

    # Pipeline 2
    c = ex["p2_counts"]
    put(ws, f"C{row}", "Step 1: Select THREE relevant and comparable pre-trained models", font=BOLD)
    put(ws, f"E{row}", "Step 2a: Test the performance of the selected models with your own dataset", font=BOLD)
    put(ws, f"C{row + 1}", "Pipeline 2 - Feature extraction (sentence embeddings for dish retrieval, not finetuned)", font=BOLD)
    put(ws, f"E{row + 2}", "Dataset:")
    put(ws, f"F{row + 2}", "data/train/p2_dev.jsonl (development), p2_test.jsonl + p2_paraphrase_test.jsonl (testing)")
    put(ws, f"E{row + 3}", "development samples (set the hybrid weight):")
    put(ws, f"F{row + 3}", c["development"])
    put(ws, f"E{row + 4}", "testing samples (literal queries):")
    put(ws, f"F{row + 4}", "=COUNTA('P2 queries'!A2:A%d)" % n2)
    put(ws, f"E{row + 5}", "testing samples (paraphrase queries):")
    put(ws, f"F{row + 5}", c["paraphrase"])
    put(ws, f"E{row + 6}", "library dishes searched:")
    put(ws, f"F{row + 6}", c["library"])
    row += 8
    header(ws, row, {"D": "Model name", "E": "Performance (top-5 accuracy, literal queries)",
                     "F": "Run time (in sec. per 100 queries, CPU)", "G": "Top-1 accuracy", "H": "Paraphrase precision@5",
                     "I": "Hybrid weight w"})
    ws.row_dimensions[row].height = 48
    row += 1
    for k, model in enumerate(ex["p2_candidates"], start=1):
        r = ex["p2_results"][model]
        col = rank_col[model]
        fill = PICK_FILL if model == ex["p2_chosen"] else None
        put(ws, f"C{row}", f"Model {k}", box=True, fill=fill)
        put(ws, f"D{row}", model, box=True, fill=fill)
        put(ws, f"E{row}", f"=COUNTIF('P2 queries'!{col}2:{col}{n2},\"<=5\")/COUNT('P2 queries'!{col}2:{col}{n2})", fmt="0.000", box=True, fill=fill)
        put(ws, f"F{row}", round(r["seconds_per_100_queries"], 2), fmt="0.00", box=True, fill=fill)
        put(ws, f"G{row}", f"=COUNTIF('P2 queries'!{col}2:{col}{n2},1)/COUNT('P2 queries'!{col}2:{col}{n2})", fmt="0.000", box=True, fill=fill)
        put(ws, f"H{row}", round(r["paraphrase_p@5"], 4), fmt="0.000", box=True, fill=fill)
        put(ws, f"I{row}", r["hybrid_weight"], box=True, fill=fill)
        row += 1
    base = "character TF-IDF (no model)"
    col = rank_col[base]
    put(ws, f"C{row}", "Baseline", box=True)
    put(ws, f"D{row}", base, box=True)
    put(ws, f"E{row}", f"=COUNTIF('P2 queries'!{col}2:{col}{n2},\"<=5\")/COUNT('P2 queries'!{col}2:{col}{n2})", fmt="0.000", box=True)
    put(ws, f"F{row}", "< 0.1", box=True)
    put(ws, f"G{row}", f"=COUNTIF('P2 queries'!{col}2:{col}{n2},1)/COUNT('P2 queries'!{col}2:{col}{n2})", fmt="0.000", box=True)
    put(ws, f"H{row}", round(ex["p2_results"][base]["paraphrase_p@5"], 4), fmt="0.000", box=True)
    row += 1
    put(ws, f"D{row}", "A query is correct when its target dish is ranked in the top 5 of the library (the menu shows 6-8 dishes). Each encoder runs "
                       "through the app's Retriever: score = cosine similarity + w x character TF-IDF, with w set on the development queries only. "
                       "Source: notebooks/Model2_experiments.ipynb; accuracy cells are formulas over sheet 'P2 queries'.", wrap=True)
    ws.merge_cells(f"D{row}:H{row}")
    ws.row_dimensions[row].height = 44
    row += 2
    put(ws, f"E{row}", "Step 2b: Pick the BEST model for implementation", font=BOLD)
    put(ws, f"E{row + 1}", f"{ex['p2_chosen']} is picked: highest top-5 accuracy; models within one standard error of the best are tied and the shortest run time wins.")

    # Overall performance (right-hand block)
    put(ws, "J8", "Step 3: Implement the App and deploy it to Streamlit Cloud", font=BOLD)
    put(ws, "J9", "1. Deployment")
    put(ws, "K10", f"App = pipeline 1 (finetuned {p1['chosen']}) + pipeline 2 ({ex['p2_chosen']}, hybrid weight {ex['p2_results'][ex['p2_chosen']]['hybrid_weight']}) + plating rules")
    put(ws, "K11", f"Pipeline 1 as deployed (finetuned model + keyword lexicon), accuracy on the hand-written sentences: {ex['p1_app_accuracy']:.3f}")
    tests = ex["app_tests"]
    put(ws, "J13", f"2. Test and evaluate the App ({len(tests)} test samples)", font=BOLD)
    put(ws, "M13", "Objective judgement: tags, dish count, no repeats, methods, soup, 20-min, balance")
    header(ws, 14, {"J": "Test", "K": "Inputs (people | request | chips)", "L": "Outputs (tags found | menu)",
                    "M": "Correct or not (notebook)", "N": "Checked on Streamlit Cloud (fill in: Correct / Incorrect)"})
    first = 15
    for i, t in enumerate(tests):
        r = first + i
        put(ws, f"J{r}", f"Test sample {i + 1}", box=True)
        put(ws, f"K{r}", f"{t['people']} people | {t['input']}", box=True, wrap=True)
        notes = f" | notes: {t['notes']}" if t["notes"] else ""
        put(ws, f"L{r}", f"tags: {t['tags found']} | {t['menu']}{notes}", box=True, wrap=True)
        put(ws, f"M{r}", t["result"] + (f" ({t['failed checks']})" if t["failed checks"] else ""), box=True, wrap=True)
        put(ws, f"N{r}", None, fill=INPUT_FILL, box=True)
        ws.row_dimensions[r].height = 62
    last = first + len(tests) - 1
    put(ws, f"L{last + 2}", "Accuracy (notebook):", font=BOLD)
    put(ws, f"M{last + 2}", f'=COUNTIF(M{first}:M{last},"Correct")&" out of "&COUNTA(M{first}:M{last})')
    put(ws, f"M{last + 3}", f'=COUNTIF(M{first}:M{last},"Correct")/COUNTA(M{first}:M{last})', fmt="0.00")
    put(ws, f"L{last + 4}", "Accuracy (Streamlit Cloud):", font=BOLD)
    put(ws, f"N{last + 4}", f'=IF(COUNTA(N{first}:N{last})=0,"not checked yet",COUNTIF(N{first}:N{last},"Correct")/COUNTA(N{first}:N{last}))', fmt="0.00")
    put(ws, f"L{last + 6}", "Legend: fill the shaded column N after repeating each request on the deployed app "
                            "(type Correct or Incorrect). Green rows = picked models.", wrap=True)
    ws.merge_cells(f"L{last + 6}:N{last + 6}")
    ws.row_dimensions[last + 6].height = 32

    for col, width in {"A": 2, "B": 2, "C": 12, "D": 50, "E": 30, "F": 24, "G": 18, "H": 22, "I": 14,
                       "J": 14, "K": 46, "L": 70, "M": 26, "N": 28}.items():
        ws.column_dimensions[col].width = width
    for r in (15,):
        ws.row_dimensions[r].height = 48
    ws.row_dimensions[14].height = 34

    out = OUT / "Experimental_results.xlsx"
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
